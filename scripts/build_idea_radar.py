# -*- coding: utf-8 -*-
"""Benchmark Idea Radar - 垂类 Benchmark 选题雷达工具"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F = 'Arial'
wb = Workbook()

# ---------- styles ----------
H_FONT = Font(name=F, bold=True, color='FFFFFF', size=11)
H_FILL = PatternFill('solid', fgColor='1F3864')
SUB_FILL = PatternFill('solid', fgColor='4472C4')
WGT_FILL = PatternFill('solid', fgColor='FFF2CC')
SEC_FONT = Font(name=F, bold=True, size=11, color='1F3864')
SEC_FILL = PatternFill('solid', fgColor='D6E4F0')
B_FONT = Font(name=F, size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=16, color='1F3864')
BLUE_IN = Font(name=F, size=10, color='0000FF')   # 手工输入
BD = Border(*[Side(style='thin', color='B4C6E7')] * 4)
WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)


def head(ws, row, n, fill=H_FILL, h=34):
    ws.row_dimensions[row].height = h
    for c in range(1, n + 1):
        x = ws.cell(row=row, column=c)
        x.font = H_FONT; x.fill = fill; x.alignment = CTR; x.border = BD


def body(ws, row, n, align=WRAP):
    for c in range(1, n + 1):
        x = ws.cell(row=row, column=c)
        x.font = B_FONT; x.alignment = align; x.border = BD


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# =====================================================================
# Sheet 1  使用说明
# =====================================================================
ws = wb.active
ws.title = '00_使用说明'
ws['A1'] = 'Benchmark Idea Radar  |  垂类 Benchmark 选题雷达'
ws['A1'].font = TITLE
ws.merge_cells('A1:D1')
ws['A2'] = '用途：系统化采集、评分、筛选垂类 Benchmark 选题，避免「做完才发现没人用 / 没区分度」'
ws['A2'].font = Font(name=F, size=10, italic=True, color='555555')
ws.merge_cells('A2:D2')

rows = [
    ['', '', '', ''],
    ['SHEET 导航', '作用', '什么时候用', '谁维护'],
    ['01_Idea池雷达', '核心工作台：录入 Idea → 8维打分 → 自动判定 GO/WATCH/NO → 自动排序；含「源类型覆盖」列检查三类源是否交叉', '每周整理 Idea 时', 'PM'],
    ['02_评分标准', '8 个维度的 1-5 分锚定描述与权重（打分时的唯一依据）', '打分前必读；季度校准', 'PM + 专家'],
    ['03_信息源雷达', '8 类 Idea 来源（S1-S8）+ 价值源/难度源/设计源 分工 + 关键信号 + 能保证的 Gate', '每日/每周扫源', 'PM'],
    ['04_空白格矩阵', '能力 × 评测范式 二维矩阵，标注饱和度，空白格＝Idea', '季度全局盘点', 'PM + 专家'],
    ['05_饱和度监测', '追踪现有 Benchmark 的 SOTA 分数，自动判定是否饱和', '月度更新', 'PM'],
    ['06_探针验收表', '20题 Pilot 的 8 条 Gate 实测验收（立项前最后一关）', 'Gate1 评审', 'PM + 算法'],
    ['07_Idea卡片', '单个 Idea 的一页纸详述模板（进入 GO 后填写）', 'Idea 转立项时', 'PM'],
    ['08_生产流水线', '★ 价值源×难度源×设计源 交叉法 6 步流程 + Failure→Capability 抽象表 + 完整示例 + 5 条纪律', 'Idea 从 0 到立项全程', 'PM'],
    ['', '', '', ''],
    ['核心方法论', '8 个源不是并列的，而是分工的', '', ''],
    ['价值源', 'S2 真实用户Query / S4 专家工作流 / S8 业务方不确定性 → 定「测什么值得测」，保证 ①真实价值 ⑧权威性', '', ''],
    ['难度源', 'S3 模型FailureCase / S7 竞品模型差异 / S6 论文与新能力 → 定「哪里有 Failure」，保证 ②Failure ③区分度', '', ''],
    ['设计源', 'S1 现有Benchmark taxonomy / S5 产品与Agent场景 → 定「怎么测」，保证 ④可客观评判 ⑤指导迭代', '', ''],
    ['交叉规则', '好 Idea 必须同时命中三类源。只有价值源→伪需求；只有难度源→无人在意的失败；只有设计源→别人 bench 的补丁', '', ''],
    ['', '', '', ''],
    ['使用流程', '', '', ''],
    ['Step 1  扫源', '按 03_信息源雷达 的频率扫描 S1-S8，原始 Idea 先粗略记入 01_Idea池', '每日 10 分钟', ''],
    ['Step 2  交叉', '按 08_生产流水线 的 6 步走：价值定位 → Failure验证 → 区分度验证 → Rubric设计 → 差异化定位', '每周', ''],
    ['Step 3  打分', '按 02_评分标准 给 8 个维度打 1-5 分，总分与判定自动生成；检查「源类型覆盖」是否达 3/3', '每周', ''],
    ['Step 4  筛选', '只看判定为「GO 立项」的行，按优先级排名取 Top 1-2', '每周', ''],
    ['Step 5  盘点', '用 04_空白格矩阵 + 05_饱和度监测 做全局校验，防止漏掉大机会', '每季度', ''],
    ['Step 6  探针', 'Top Idea 做 20 题 Pilot，填 06_探针验收表，8 条 Gate 全过才立项', '立项前', ''],
    ['Step 7  立项', '填 07_Idea卡片 作为 PRD 前置文档', '立项时', ''],
    ['', '', '', ''],
    ['判定门槛', '', '', ''],
    ['GO   立项', '加权总分 >= 4.00  且  4 个核心维度均 >= 3 分', '', ''],
    ['WATCH 观察', '加权总分 3.20 ~ 3.99  且无红线', '进 Idea 池孵化，等条件成熟', ''],
    ['NO   否决', '加权总分 < 3.20  或  触发红线', '', ''],
    ['红线规则', '① 真实任务价值 / ② 明显Failure / ③ 区分度潜力 / ④ 可客观评判 —— 任一 <= 2 分即否决', '核心四项不可妥协', ''],
    ['', '', '', ''],
    ['颜色约定', '', '', ''],
    ['蓝色字体', '需要手工输入的单元格', '', ''],
    ['黑色字体', '公式自动计算，请勿覆盖', '', ''],
    ['黄色底纹', '权重行 / 关键假设，可按团队策略调整', '', ''],
]
for i, r in enumerate(rows, 3):
    ws.append(r)
    if r[0] in ('SHEET 导航',):
        head(ws, i, 4)
    elif r[0] in ('使用流程', '判定门槛', '颜色约定'):
        for c in range(1, 5):
            x = ws.cell(row=i, column=c); x.font = SEC_FONT; x.fill = SEC_FILL; x.border = BD; x.alignment = WRAP
    elif any(r):
        body(ws, i, 4)
        ws.cell(row=i, column=1).font = BOLD
widths(ws, {'A': 20, 'B': 78, 'C': 34, 'D': 14})

# =====================================================================
# Sheet 2  Idea 池雷达   (核心)
# =====================================================================
p = wb.create_sheet('01_Idea池雷达')
DIMS = ['①真实任务价值', '②明显Failure', '③区分度潜力', '④可客观评判',
        '⑤指导迭代', '⑥抗污染寿命', '⑦成本可行', '⑧权威性可采信']
WEIGHTS = [0.20, 0.20, 0.15, 0.15, 0.10, 0.08, 0.07, 0.05]

hdr = (['编号', 'Idea 名称', '垂类领域', '目标能力', '评测范式', '来源象限', '来源渠道', '一句话描述']
       + DIMS + ['加权总分', '红线', '判定', '优先级', '已有竞品 Benchmark', '下一步动作', '负责人', '状态'])
p.append(hdr)
head(p, 1, len(hdr), h=46)

# 权重行
wrow = ['权重 →', '', '', '', '', '', '', ''] + WEIGHTS + ['= 100%', '', '', '', '', '', '', '']
p.append(wrow)
for c in range(1, len(hdr) + 1):
    x = p.cell(row=2, column=c); x.font = BOLD; x.fill = WGT_FILL; x.alignment = CTR; x.border = BD
for c in range(9, 17):
    p.cell(row=2, column=c).number_format = '0%'

IDEAS = [
    ['ID-01', '用药安全 Rubric 评测（剂量/相互作用/禁忌）', '医疗', '用药安全决策', 'Rubric 开放式',
     '③需求驱动', '甲方访谈 + BadCase', '真实处方场景，逐点核查剂量、相互作用、禁忌、肾功能调整，每点绑定说明书/指南',
     5, 5, 4, 4, 5, 4, 3, 4, '', '', '', '', 'MedBench(仅少量选择题)', '出 20 题探针 + 招募临床药师', 'PM-A', '探针中'],
    ['ID-02', '指南溯源循证推理评测（S/A级依据链）', '医疗', '循证临床推理', 'Rubric + 依据溯源',
     '②方法驱动', 'arXiv + 自研', '开放式病例问答，每个评分点必须绑定 S/A 级权威指南，形成可审计依据链',
     5, 4, 4, 5, 5, 4, 2, 5, '', '', '', '', 'ClinConsensus / LiveMedBench', '完成 10 题标注，做 IAA 一致性检验', 'PM-A', '进行中'],
    ['ID-03', '多轮问诊主动提问能力评测', '医疗', '信息采集 / 主动澄清', '多轮交互',
     '②方法驱动', 'AMIE 论文 + OSCE', '模拟患者 Agent，评估模型是否会主动追问关键鉴别信息而非直接下结论',
     5, 5, 4, 3, 4, 5, 2, 3, '', '', '', '', 'AMIE(未开源) / MedQA-CS', '先解决患者模拟器保真度问题', 'PM-B', '设计中'],
    ['ID-04', '长病历长上下文理解与信息抽取', '医疗', '长上下文 / 信息抽取', '结构化抽取',
     '⑧新能力', '榜单空白 + 客户需求', '10万字级真实住院病历，考察跨页信息整合、时间线重建、矛盾信息识别',
     4, 4, 4, 4, 4, 4, 3, 3, '', '', '', '', '几乎空白', '解决病历脱敏与伦理审批', 'PM-B', '待启动'],
    ['ID-05', '医疗安全边界与拒答能力评测', '医疗', '安全 / 拒答边界', 'Rubric + 分类',
     '④合规驱动', '监管政策 + BadCase', '高危问询（自杀、超说明书用药、诊断替代）下是否恰当拒答并引导就医',
     5, 4, 3, 4, 4, 4, 4, 4, '', '', '', '', 'TRIDENT-Bench(偏通用)', '对齐监管口径后出题', 'PM-A', '待启动'],
    ['ID-06', '医学影像+文本联合推理评测', '医疗', '多模态推理', '多模态 QA',
     '②方法驱动', 'MedXpertQA + 顶会', '影像 + 病史联合输入，考察影像描述到诊断结论的推理链正确性',
     4, 5, 4, 3, 3, 4, 2, 4, '', '', '', '', 'MedXpertQA / OmniMedVQA', '影像版权与标注成本待评估', 'PM-C', '观察'],
    ['ID-07', '法律意见书质量 Rubric 评测', '法律', '法律文书生成 / 论证', 'Rubric 开放式',
     '②方法驱动', '范式迁移（医疗→法律）', '真实咨询场景出具法律意见书，逐点核查法条引用、争议识别、风险提示、结论明确性',
     5, 4, 4, 4, 4, 4, 3, 4, '', '', '', '', 'LawBench(20任务,偏抽取分类)', '招募执业律师标注团队', 'PM-C', '探针中'],
    ['ID-08', '合同起草与审查真实任务（SWE-bench 式）', '法律', '合同起草 / 风险审查', '真实任务 + 可执行检查',
     '②方法驱动', '范式迁移（代码→法律）', '给定交易背景起草/修订合同，用条款清单与风险规则做半自动判定',
     5, 4, 4, 3, 4, 4, 2, 3, '', '', '', '', '空白', '设计条款级自动检查规则', 'PM-C', '设计中'],
    ['ID-09', '投研报告生成质量 Rubric 评测', '金融', '投研分析 / 观点论证', 'Rubric 开放式',
     '③需求驱动', '甲方访谈', '给定财报与行业数据生成投研观点，核查数据准确性、逻辑链、风险披露、结论可执行性',
     4, 4, 4, 3, 4, 3, 3, 3, '', '', '', '', 'CFBenchmark / CNFinBench', '解决「多解性」评分难题', 'PM-D', '观察'],
    ['ID-10', '财报长上下文数值推理评测', '金融', '数值推理 / 长上下文', '精确答案 + 过程分',
     '①饱和驱动', '榜单饱和信号', '多期财报交叉计算（同比/调整后口径/附注勾稽），答案唯一可自动判分',
     5, 4, 5, 5, 4, 4, 4, 3, '', '', '', '', 'FinanceBench(规模小)', '直接扩题量，最快可交付', 'PM-D', '探针中'],
    ['ID-11', 'CPA 审计程序推理评测', '审计', '审计程序 / 职业判断', 'Rubric 开放式',
     '⑤专业体系', 'CPA 科目体系映射', '给定被审计单位情形，设计审计程序并判断风险，映射 CPA 与审计准则条文',
     4, 5, 4, 4, 4, 5, 3, 3, '', '', '', '', '基本空白', '锁定准则版本，招募注会专家', 'PM-D', '待启动'],
    ['ID-12', '垂类 Agent 工具调用能力评测（医疗助手）', '医疗', 'Agent / 工具调用', 'Agent 任务 + 轨迹评分',
     '⑧新能力', '新能力空白', '给定知识库/计算器/检索工具，评估是否正确调用、参数是否准确、结论是否落地',
     4, 5, 5, 4, 4, 5, 2, 3, '', '', '', '', '垂类 Agent 评测几乎空白', '先搭工具沙箱环境', 'PM-B', '设计中'],
    ['ID-13', '垂类防污染动态更新评测框架（方法论）', '跨领域', '评测基建 / 抗污染', '框架 + 管道',
     '②方法驱动', 'LiveMedBench + 榜单异常', '时间切片 + 私有 holdout + 等价变体生成，做成可复用的抗污染基建',
     3, 3, 3, 4, 3, 5, 3, 4, '', '', '', '', 'LiveMedBench', '作为底层能力嵌入其他项目', 'PM-A', '孵化'],
    ['ID-14', 'K12 解题过程评分评测', '教育', '解题过程 / 步骤分', 'Rubric 步骤分',
     '③需求驱动', '客户需求', '数学题按步骤给分，考察过程正确性而非仅答案',
     3, 2, 2, 3, 3, 2, 4, 2, '', '', '', '', 'EduBench 等已较多', '现有模型已基本满分，无区分度', 'PM-E', '否决'],
    ['ID-15', '罕见病分子机制知识问答', '医疗', '知识记忆', '选择题',
     '①饱和驱动', 'arXiv', '考察罕见病致病基因与分子通路的知识记忆',
     1, 2, 2, 5, 2, 2, 4, 2, '', '', '', '', 'MedQA / CMB 已覆盖', '真实工作流不做此任务，否决', 'PM-E', '否决'],
    ['ID-16', 'PDF 财报异常识别与证据引用（任务式）', '金融', '长上下文+数值接地+跨页推理+引用正确性', '真实任务 + 自动判分',
     '①饱和驱动', '范式迁移', '给真实年报 PDF：读财报→找异常→引用证据页码→给结论。答案唯一可自动判分，citation 正确性单独计分',
     5, 5, 5, 4, 5, 4, 4, 4, '', '', '', '', 'FinanceBench(规模小) / CNFinBench(偏单轮QA)', '★流水线完整走通示例：见 08_生产流水线', 'PM-D', '立项'],
]

start = 3
for i, row in enumerate(IDEAS):
    r = start + i
    p.append(row)
    body(p, r, len(hdr))
    p.row_dimensions[r].height = 50
    # 手工输入区标蓝
    for c in list(range(1, 9)) + list(range(9, 17)) + [21, 22, 23, 24]:
        p.cell(row=r, column=c).font = BLUE_IN
    for c in range(9, 17):
        p.cell(row=r, column=c).alignment = CTR
    # 公式
    p.cell(row=r, column=17).value = f'=SUMPRODUCT($I$2:$P$2,I{r}:P{r})'
    p.cell(row=r, column=17).number_format = '0.00'
    p.cell(row=r, column=17).font = Font(name=F, bold=True, size=11)
    p.cell(row=r, column=17).alignment = CTR
    p.cell(row=r, column=18).value = f'=IF(MIN(I{r}:L{r})<=2,"⛔红线","✓")'
    p.cell(row=r, column=18).alignment = CTR
    p.cell(row=r, column=19).value = (
        f'=IF(R{r}="⛔红线","NO 否决",'
        f'IF(Q{r}>=4,"GO 立项",IF(Q{r}>=3.2,"WATCH 观察","NO 否决")))')
    p.cell(row=r, column=19).alignment = CTR
    p.cell(row=r, column=19).font = Font(name=F, bold=True, size=10)
    p.cell(row=r, column=20).value = (
        f'=IF(S{r}="NO 否决","-",RANK(Q{r},$Q${start}:$Q${start+len(IDEAS)-1}))')
    p.cell(row=r, column=20).alignment = CTR

last = start + len(IDEAS) - 1

# 来源列改写为 S 码（对应 03_信息源雷达），并标注三类源交叉数
SRC_MAP = ['S8+S3', 'S6+S1', 'S6+S4', 'S5+S8', 'S8+S3', 'S6+S1', 'S1+S4', 'S1+S5',
           'S2+S4', 'S1+S2', 'S4+S1', 'S5+S3', 'S1+S6', 'S2', 'S6', 'S8+S3+S5']
VALUE_S = {'S2', 'S4', 'S8'}
DIFF_S = {'S3', 'S6', 'S7'}
DESIGN_S = {'S1', 'S5'}
p.cell(row=1, column=6).value = '来源 S 码'
p.cell(row=1, column=7).value = '源类型覆盖\n(价值/难度/设计)'
for i in range(len(IDEAS)):
    r = start + i
    code = SRC_MAP[i]
    p.cell(row=r, column=6).value = code
    p.cell(row=r, column=6).alignment = CTR
    p.cell(row=r, column=6).font = BLUE_IN
    parts = set(code.split('+'))
    tags = []
    if parts & VALUE_S:
        tags.append('价值')
    if parts & DIFF_S:
        tags.append('难度')
    if parts & DESIGN_S:
        tags.append('设计')
    cov = '+'.join(tags) + (f'  ({len(tags)}/3)' if tags else '')
    x = p.cell(row=r, column=7, value=cov)
    x.alignment = CTR
    x.border = BD
    if len(tags) == 3:
        x.font = Font(name=F, bold=True, size=10, color='FFFFFF')
        x.fill = PatternFill('solid', fgColor='00B050')
    elif len(tags) == 2:
        x.font = Font(name=F, bold=True, size=10)
        x.fill = PatternFill('solid', fgColor='FFEB9C')
    else:
        x.font = Font(name=F, size=10, color='C00000')
        x.fill = PatternFill('solid', fgColor='FFC7CE')

# 数据校验：1-5 分
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5',
                    allow_blank=True, showErrorMessage=True,
                    errorTitle='打分范围', error='请输入 1 到 5 的整数')
p.add_data_validation(dv)
dv.add(f'I{start}:P{last}')

# 条件格式
rng = f'I{start}:P{last}'
p.conditional_formatting.add(rng, CellIsRule(operator='lessThanOrEqual', formula=['2'],
    fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))
p.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['3'],
    fill=PatternFill('solid', bgColor='FFEB9C'), font=Font(color='9C5700')))
p.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['4'],
    fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100')))

p.conditional_formatting.add(f'S{start}:S{last}', FormulaRule(
    formula=[f'ISNUMBER(SEARCH("GO",S{start}))'],
    fill=PatternFill('solid', bgColor='00B050'), font=Font(color='FFFFFF', bold=True)))
p.conditional_formatting.add(f'S{start}:S{last}', FormulaRule(
    formula=[f'ISNUMBER(SEARCH("WATCH",S{start}))'],
    fill=PatternFill('solid', bgColor='FFC000'), font=Font(color='000000', bold=True)))
p.conditional_formatting.add(f'S{start}:S{last}', FormulaRule(
    formula=[f'ISNUMBER(SEARCH("NO",S{start}))'],
    fill=PatternFill('solid', bgColor='C00000'), font=Font(color='FFFFFF', bold=True)))
p.conditional_formatting.add(f'R{start}:R{last}', FormulaRule(
    formula=[f'R{start}="⛔红线"'],
    fill=PatternFill('solid', bgColor='C00000'), font=Font(color='FFFFFF', bold=True)))
p.conditional_formatting.add(f'Q{start}:Q{last}', CellIsRule(
    operator='greaterThanOrEqual', formula=['4'],
    fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100', bold=True)))

# 汇总行
sm = last + 2
p.cell(row=sm, column=1, value='汇总').font = SEC_FONT
p.cell(row=sm, column=1).fill = SEC_FILL
p.cell(row=sm, column=2, value='GO 立项数').font = BOLD
p.cell(row=sm, column=3).value = f'=COUNTIF(S{start}:S{last},"GO 立项")'
p.cell(row=sm, column=4, value='WATCH 观察数').font = BOLD
p.cell(row=sm, column=5).value = f'=COUNTIF(S{start}:S{last},"WATCH 观察")'
p.cell(row=sm, column=6, value='NO 否决数').font = BOLD
p.cell(row=sm, column=7).value = f'=COUNTIF(S{start}:S{last},"NO 否决")'
p.cell(row=sm, column=8, value='平均总分').font = BOLD
p.cell(row=sm, column=9).value = f'=ROUND(AVERAGE(Q{start}:Q{last}),2)'
for c in range(1, 10):
    p.cell(row=sm, column=c).border = BD
    if c in (3, 5, 7, 9):
        p.cell(row=sm, column=c).font = Font(name=F, bold=True, size=11, color='C00000')
        p.cell(row=sm, column=c).alignment = CTR

widths(p, {'A': 8, 'B': 34, 'C': 10, 'D': 16, 'E': 18, 'F': 13, 'G': 20, 'H': 46,
           'I': 9, 'J': 9, 'K': 9, 'L': 9, 'M': 9, 'N': 9, 'O': 9, 'P': 9,
           'Q': 10, 'R': 8, 'S': 12, 'T': 8, 'U': 26, 'V': 30, 'W': 9, 'X': 10})
p.freeze_panes = 'C3'
p.auto_filter.ref = f'A1:X{last}'

# =====================================================================
# Sheet 3  评分标准
# =====================================================================
s = wb.create_sheet('02_评分标准')
s['A1'] = '8 维评分锚定标准（打分唯一依据）'
s['A1'].font = TITLE
s.merge_cells('A1:H1')
hdr3 = ['维度', '核心问题', '权重', '1 分', '2 分', '3 分', '4 分', '5 分']
s.append([])
s.append(hdr3)
head(s, 3, 8)

SC = [
    ['①真实任务价值', '真实工作流里有人做吗？做错有代价吗？', 0.20,
     '纯学术/考试题，真实工作流不做', '偶尔涉及，错了无实质后果', '有人做，但非核心环节',
     '高频任务，做错有明确损失', '核心高频任务，做错有严重后果（安全/资金/法律）'],
    ['②明显Failure', 'SOTA 模型确实做不好吗？', 0.20,
     'SOTA 已接近满分（>90%）', 'SOTA 表现良好（80-90%）', 'SOTA 中等（70-80%）',
     'SOTA 明显吃力（50-70%）', 'SOTA 大面积失败（30-50%），且失败有规律'],
    ['③区分度潜力', '强弱模型能被分开吗？', 0.15,
     '预期所有模型分数接近', '仅能区分强弱两档', '能分 3 档，但排序不稳',
     '分数梯度明显，排序稳定', '梯度清晰且可随模型进步持续区分（有难度分层设计）'],
    ['④可客观评判', '同一答案重复评分结果一致吗？', 0.15,
     '完全主观，无法复现', '需专家主观裁量，一致性 <0.5', '可写 Rubric，一致性 0.5-0.7',
     'Rubric 二元可判定，一致性 0.7-0.85', '答案唯一或 Rubric 完全可机判，一致性 >0.85'],
    ['⑤指导迭代', '模型得 45 分，研发知道该改什么吗？', 0.10,
     '只能给一个总分', '能分大类但无法定位', '能定位到能力维度',
     '能定位到子能力 + 错误类型', '能定位到子能力、错误类型并给出数据/训练建议'],
    ['⑥抗污染寿命', '半年后还有效吗？', 0.08,
     '题目全在公网，已被训练', '大部分可检索到', '部分可检索，无更新机制',
     '有私有 holdout 或时间切片', '原创题 + 私有集 + 可持续更新管道，寿命 >2 年'],
    ['⑦成本可行', '数据/专家/预算够吗？', 0.07,
     '数据拿不到或有版权死结', '专家极难招募，成本远超预算', '成本紧张，需砍规模',
     '成本可控，资源基本齐备', '数据可持续获取，专家到位，成本充裕'],
    ['⑧权威性可采信', '有人会认你的结论吗？', 0.05,
     '无背书，方法不公开', '仅内部认可', '有部分专家参与',
     '专家资质可公示 + 依据可溯源', '权威专家背书 + 依据全溯源 + 方法论可发表复现'],
]
for i, r in enumerate(SC, 4):
    s.append(r)
    body(s, i, 8)
    s.row_dimensions[i].height = 56
    s.cell(row=i, column=1).font = SEC_FONT
    s.cell(row=i, column=1).fill = SEC_FILL
    s.cell(row=i, column=3).number_format = '0%'
    s.cell(row=i, column=3).alignment = CTR
    s.cell(row=i, column=3).font = Font(name=F, bold=True, size=10)
    s.cell(row=i, column=3).fill = WGT_FILL

r = len(SC) + 4
s.cell(row=r, column=1, value='权重合计').font = BOLD
s.cell(row=r, column=3).value = '=SUM(C4:C11)'
s.cell(row=r, column=3).number_format = '0%'
s.cell(row=r, column=3).font = Font(name=F, bold=True, size=11, color='C00000')
s.cell(row=r, column=3).alignment = CTR
for c in range(1, 9):
    s.cell(row=r, column=c).border = BD

r += 2
notes = [
    ['红线规则', '① ② ③ ④ 四个核心维度任一 <= 2 分 → 直接否决，不看总分'],
    ['判定门槛', 'GO >= 4.00 ；WATCH 3.20-3.99 ；NO < 3.20'],
    ['权重可调', '权重行是团队策略的体现：偏商业化可提高①③；偏学术可提高②⑧；01 表会自动重算'],
    ['打分纪律', '打分必须给出「凭什么给这个分」的一句话理由，写在 01 表「下一步动作」或备注里'],
    ['②的反直觉点', 'SOTA 分数过低（<20%）不加分反而扣分——通常说明题目设计或评分标准有问题，而非模型无能'],
    ['③与②的张力', '若所有模型都接近 0 分：有 Failure 但无区分度。必须做难度分层（易30%/中40%/难30%）'],
    ['①与④的张力', '越真实的任务越开放、越难客观评分。解法是 case-specific Rubric + 依据溯源，代价是成本'],
]
s.cell(row=r - 1, column=1, value='关键规则与陷阱').font = SEC_FONT
s.cell(row=r - 1, column=1).fill = SEC_FILL
for c in range(1, 9):
    s.cell(row=r - 1, column=c).fill = SEC_FILL
    s.cell(row=r - 1, column=c).border = BD
for i, (k, v) in enumerate(notes):
    rr = r + i
    s.cell(row=rr, column=1, value=k).font = BOLD
    s.cell(row=rr, column=2, value=v).font = B_FONT
    s.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
    for c in range(1, 9):
        s.cell(row=rr, column=c).border = BD
        s.cell(row=rr, column=c).alignment = WRAP
    s.row_dimensions[rr].height = 30

widths(s, {'A': 18, 'B': 34, 'C': 8, 'D': 28, 'E': 28, 'F': 28, 'G': 28, 'H': 34})

# =====================================================================
# Sheet 4  信息源雷达
# =====================================================================
src = wb.create_sheet('03_信息源雷达')
src['A1'] = 'Idea 来源雷达：8 类源 × 分工 × 监测节奏'
src['A1'].font = TITLE
src.merge_cells('A1:I1')
src['A2'] = ('核心原则：8 个源不是并列的，而是分工的 —— 价值源定「测什么值得测」，难度源定「哪里有 Failure」，'
             '设计源定「怎么测」。好 Idea = 价值源 × 难度源 × 设计源 三者交叉。')
src['A2'].font = Font(name=F, size=10, italic=True, color='C00000')
src.merge_cells('A2:I2')
src.append([])
hdr4 = ['#', '来源', '分工类型', '具体信息源 / 代表', '监测什么（重点不是照搬题）',
        '关键信号 → 立即记 Idea', '能保证的 Gate', '频率', '负责人 / 上次扫描']
src.append(hdr4)
head(src, 4, 9, h=40)

VAL = PatternFill('solid', fgColor='E2EFDA')   # 价值源
DIF = PatternFill('solid', fgColor='FCE4D6')   # 难度源
DES = PatternFill('solid', fgColor='DDEBF7')   # 设计源

SRC = [
    ['S1', '现有 Benchmark\n/ Leaderboard', '设计源',
     'GAIA / Humanity\'s Last Exam / BrowseComp / SWE-bench Verified / τ²-bench\nMedBench·LawBench·CNFinBench·CFBenchmark·FinanceBench\nOpenCompass 司南 / LMArena / Vals AI',
     '① taxonomy 怎么切能力\n② task design 怎么把能力变成可判分任务\n③ 它自己披露的 failure case\n④ SOTA 分数与饱和度\n（重点不是照搬题）',
     '• Top 模型 >90% 且密集 → 已饱和\n• taxonomy 有明显缺项 → 覆盖不深\n• 只有单轮 QA → 可升级为任务式\n• 半年增幅 >15pt → 疑似污染',
     '④可客观评判\n⑤指导迭代\n差异化定位',
     '每周', 'PM'],
    ['S2', '真实用户 Query', '价值源 ⭐⭐⭐',
     '线上请求日志 / 客服与工单记录 / 用户反馈 / 搜索词 / 甲方提供的真实场景样本',
     'Query 聚类 → 找「高频 × 高价值 × 当前模型表现差」的簇\n注意：用户要的往往不是考试题',
     '• 高频簇 + 模型表现差 → 头号候选\n• 例：金融用户要的不是 CFA 选择题，\n  而是「读财报→找异常→引用证据→给结论」\n• 用户反复换问法 = 模型没答好',
     '①真实任务价值（最强）\n明确使用者与场景',
     '每周聚类', 'PM + 数据'],
    ['S3', '模型 Failure Case', '难度源 ⭐⭐⭐\n最值得持续挖',
     '拿当前 SOTA 跑真实任务采集失败 / 内部 BadCase 库 / 专家挑错记录 / 线上负反馈',
     '把失败案例沉淀 → 归类失败模式 → 抽象成可评测的能力\n（Failure → Capability 的抽象是关键动作）',
     '• 同类错误反复出现 → 能力缺口\n• 例：长 PDF 总引用错数字 → 抽象为\n  Long-context retrieval + Numerical grounding\n  + Cross-page reasoning + Citation correctness',
     '②明显Failure（最强）\n③区分度',
     '每周归类', 'PM + 算法'],
    ['S4', '专家工作流', '价值源 ⭐⭐⭐',
     '律师 / 医生 / 金融分析师 / 科研人员 / 注会 的真实日常工作',
     '研究他们每天到底在完成什么工作、卡在哪里\n⛔ 反模式：问专家「给我出 100 道难题」',
     '✅ 专家三问法（必背）：\n• 你工作中最难的 5 件事是什么？\n• 新人最容易在哪犯错？\n• 什么任务你一看就知道对方专不专业？\n→ 这三问的回答直接变成 Rubric 核心项',
     '①真实任务价值\n⑧权威性可采信\nRubric 设计依据',
     '每月 1-2 场', 'PM + 专家'],
    ['S5', '产品 / Agent 场景', '设计源',
     'Deep Research / Browser Agent / Computer Use / Coding Agent / 搜索与检索增强',
     '模型正在从「回答问题」走向哪些任务形态\n→ 垂类评测同步从单轮 QA 升级为「完成一个专业任务」',
     '• 新任务形态出现 → 垂类版本几乎必然空白\n• 单轮 QA 的垂类 bench → 可升级为任务式\n• 工具调用 / 多步规划 / 轨迹评分 是蓝海',
     '④可客观评判（任务可判定）\n时效性与前瞻性',
     '每月', 'PM'],
    ['S6', '论文与新能力', '难度源',
     'arXiv 关键词订阅 / NeurIPS D&B Track / ICLR / ACL·EMNLP / NEJM AI / Nature Digital Medicine',
     '看论文不为学算法，为找 evaluation idea\n可与垂类结合的新能力维度',
     '• hallucination / uncertainty calibration\n• long context / multimodal reasoning\n• tool use / citation correctness\n• self-correction / agent planning\n• 论文 Limitations 章节 ← 最高价值来源',
     '②明显Failure\n⚠ ①真实价值需另行验证',
     '每日 10 分钟', 'PM'],
    ['S7', '竞品模型差异', '难度源 ⭐⭐\n最便宜的验证器',
     '同一批任务跑 GPT / Claude / Gemini / DeepSeek / Qwen 等',
     '看模型之间在哪类任务上差距特别大\n（20 题 × 4 模型即可，成本极低）',
     '• 模型间差距大 → 优质 Benchmark 候选\n• 无区分度的题通常价值有限 → 直接否决\n• 排序与公认能力一致 → 信号可信\n• 排序混乱 → 是噪声不是区分度',
     '③区分度（直接验证）',
     '每次有新模型\n+ 每轮探针', 'PM + 算法'],
    ['S8', '业务方反馈\n（不确定性）', '价值源 ⭐⭐⭐',
     '模型团队 / 产品团队 / 甲方决策者 / 招投标技术评分表',
     '核心问法：「现在最不知道模型哪项能力到底行不行？」\n这种不确定性本身就是 Benchmark Idea',
     '• 例：团队不知道新模型有没有提升 PDF 财务分析\n  → 专门设计这一能力集\n• 「上次选模型最纠结什么」\n• 招标技术评分项 = 现成能力清单\n• 厂商集体回避的能力 = 评测空白',
     '①真实价值中的「谁用」\n⑤指导迭代（需求方明确）',
     '每两周同步', 'PM'],
]
for i, r in enumerate(SRC, 5):
    src.append(r)
    body(src, i, 9)
    src.row_dimensions[i].height = 108
    src.cell(row=i, column=1).alignment = CTR
    src.cell(row=i, column=1).font = BOLD
    src.cell(row=i, column=2).font = SEC_FONT
    src.cell(row=i, column=2).fill = SEC_FILL
    src.cell(row=i, column=2).alignment = CTR
    t = r[2]
    fill = VAL if '价值源' in t else (DIF if '难度源' in t else DES)
    src.cell(row=i, column=3).fill = fill
    src.cell(row=i, column=3).font = Font(name=F, bold=True, size=10)
    src.cell(row=i, column=3).alignment = CTR
    src.cell(row=i, column=7).alignment = CTR
    src.cell(row=i, column=8).alignment = CTR
    src.cell(row=i, column=9).font = BLUE_IN

r = 5 + len(SRC) + 1
src.cell(row=r, column=1, value='分工说明：三类源各自的能力与致命缺陷').font = SEC_FONT
for c in range(1, 10):
    src.cell(row=r, column=c).fill = SEC_FILL
    src.cell(row=r, column=c).border = BD
src.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)

CLS = [
    ['价值源', 'S2 真实用户Query / S4 专家工作流 / S8 业务方不确定性',
     '保证 ①真实任务价值、⑧权威性，明确使用者与决策场景',
     '⚠ 单独用的缺陷：不知道模型是否真做不好 → 可能做完发现已饱和', VAL],
    ['难度源', 'S3 Failure Case / S7 竞品模型差异 / S6 论文与新能力',
     '保证 ②明显Failure、③区分度',
     '⚠ 单独用的缺陷：不知道该 failure 在真实工作流里重不重要 → 伪需求', DIF],
    ['设计源', 'S1 现有Benchmark taxonomy / S5 产品与Agent场景',
     '保证 ④可客观评判、⑤指导迭代，确定差异化定位',
     '⚠ 单独用的缺陷：容易变成别人 benchmark 的补丁，缺原生价值', DES],
]
for i, (k, v1, v2, v3, fl) in enumerate(CLS, r + 1):
    src.cell(row=i, column=1, value=k).font = Font(name=F, bold=True, size=10)
    src.cell(row=i, column=1).fill = fl
    src.cell(row=i, column=1).alignment = CTR
    src.cell(row=i, column=2, value=v1).font = B_FONT
    src.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    src.cell(row=i, column=5, value=v2).font = B_FONT
    src.merge_cells(start_row=i, start_column=5, end_row=i, end_column=6)
    src.cell(row=i, column=7, value=v3).font = Font(name=F, size=10, color='C00000')
    src.merge_cells(start_row=i, start_column=7, end_row=i, end_column=9)
    for c in range(1, 10):
        src.cell(row=i, column=c).border = BD
        src.cell(row=i, column=c).alignment = WRAP
    src.row_dimensions[i].height = 40

widths(src, {'A': 6, 'B': 18, 'C': 15, 'D': 44, 'E': 40, 'F': 52, 'G': 24, 'H': 14, 'I': 14})
src.freeze_panes = 'A5'

# =====================================================================
# Sheet 5  空白格矩阵
# =====================================================================
g = wb.create_sheet('04_空白格矩阵')
g['A1'] = '能力 × 评测范式 空白格矩阵（示例：医疗垂类）'
g['A1'].font = TITLE
g.merge_cells('A1:H1')
g['A2'] = '用法：把你负责的垂类画成此表 → 标注饱和度 → 空白格即 Idea → 送入 01_Idea池 打分'
g['A2'].font = Font(name=F, size=10, italic=True, color='555555')
g.merge_cells('A2:H2')

hdr5 = ['能力维度 \\ 评测范式', '选择题', '开放问答', 'Rubric 评分', '多轮交互', 'Agent 任务', '多模态', '空白格机会说明']
g.append([])
g.append(hdr5)
head(g, 4, 8)

GAP = [
    ['知识记忆', '●饱和', '●饱和', '—', '—', '—', '○少', '无机会，勿入'],
    ['鉴别诊断', '●饱和', '◐较多', '★在做', '○少', '□空白', '○少', 'Agent 化鉴别诊断为空白'],
    ['治疗方案制定', '◐较多', '○少', '★缺口', '□空白', '□空白', '—', 'Rubric 化治疗方案评测是明确缺口'],
    ['用药安全', '○少', '□空白', '★★大缺口', '□空白', '□空白', '—', '最高优先级缺口（ID-01）'],
    ['指南溯源 / 循证', '□空白', '□空白', '★★优势区', '—', '□空白', '—', '自研优势方向（ID-02）'],
    ['长病历理解', '□空白', '○少', '○少', '—', '□空白', '□空白', '长上下文垂类评测普遍空白（ID-04）'],
    ['医患沟通', '—', '○少', '○少', '★缺口', '□空白', '—', '多轮交互评测最空白（ID-03）'],
    ['安全 / 拒答边界', '○少', '○少', '★缺口', '□空白', '□空白', '—', '合规驱动刚需（ID-05）'],
    ['工具调用 / 计算', '□空白', '□空白', '□空白', '□空白', '★★大缺口', '—', '垂类 Agent 评测蓝海（ID-12）'],
]
for i, r in enumerate(GAP, 5):
    g.append(r)
    body(g, i, 8)
    g.row_dimensions[i].height = 26
    g.cell(row=i, column=1).font = BOLD
    g.cell(row=i, column=1).fill = SEC_FILL
    for c in range(2, 8):
        g.cell(row=i, column=c).alignment = CTR
        v = str(g.cell(row=i, column=c).value or '')
        if '★★' in v:
            g.cell(row=i, column=c).fill = PatternFill('solid', fgColor='FF0000')
            g.cell(row=i, column=c).font = Font(name=F, bold=True, size=10, color='FFFFFF')
        elif '★' in v:
            g.cell(row=i, column=c).fill = PatternFill('solid', fgColor='FFC000')
            g.cell(row=i, column=c).font = Font(name=F, bold=True, size=10)
        elif '□空白' in v:
            g.cell(row=i, column=c).fill = PatternFill('solid', fgColor='DDEBF7')
        elif '●饱和' in v:
            g.cell(row=i, column=c).fill = PatternFill('solid', fgColor='D9D9D9')
            g.cell(row=i, column=c).font = Font(name=F, size=10, color='808080')

lg = len(GAP) + 6
g.cell(row=lg, column=1, value='图例').font = SEC_FONT
g.cell(row=lg, column=1).fill = SEC_FILL
LEG = [
    ['●饱和', '已有成熟 benchmark 且 SOTA 高分 → 无机会'],
    ['◐较多', '已有多个 benchmark，需找差异化角度'],
    ['○少', '仅有零星工作 → 有机会'],
    ['□空白', '基本无人做 → 潜在机会，但需验证真实价值'],
    ['★缺口', '空白 + 高真实价值 → 优先候选'],
    ['★★大缺口', '空白 + 高价值 + 有 Failure → 立即立项候选'],
    ['—', '该组合不适用'],
]
for i, (k, v) in enumerate(LEG, lg + 1):
    g.cell(row=i, column=1, value=k).font = BOLD
    g.cell(row=i, column=2, value=v).font = B_FONT
    g.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    for c in range(1, 9):
        g.cell(row=i, column=c).border = BD
        g.cell(row=i, column=c).alignment = WRAP
widths(g, {'A': 22, 'B': 12, 'C': 12, 'D': 15, 'E': 13, 'F': 14, 'G': 12, 'H': 40})

# =====================================================================
# Sheet 6  饱和度监测
# =====================================================================
sat = wb.create_sheet('05_饱和度监测')
sat['A1'] = '现有 Benchmark 饱和度监测（月度更新 SOTA 分数）'
sat['A1'].font = TITLE
sat.merge_cells('A1:J1')
sat.append([])
hdr6 = ['Benchmark', '垂类', '题型', '规模', 'SOTA 得分%', '半年前得分%', '增幅', '饱和判定', '上次更新', '衍生 Idea 机会']
sat.append(hdr6)
head(sat, 3, 10)

SAT = [
    ['MedQA (USMLE)', '医疗', '选择题', '~12K', 93, 91, None, None, '2020', '已饱和 → 转开放式/Rubric'],
    ['MedQA (MCMLE)', '医疗', '选择题', '~34K', 89, 85, None, None, '2020', '接近饱和 → 提难度'],
    ['CMB', '医疗', '选择题', '280K+', 86, 80, None, None, '2023', '规模大但题型单一 → 加推理链评测'],
    ['MedBench', '医疗', '综合', '多任务', 78, 71, None, None, '2026', '安全伦理维度仍偏低 → 专项深挖'],
    ['MedXpertQA', '医疗', '选择+多模态', '—', 58, 44, None, None, '2025', '难度高，仍有空间'],
    ['LawBench', '法律', '20任务', '—', 74, 68, None, None, '2023', '偏抽取分类 → 补文书生成 Rubric'],
    ['CFBenchmark', '金融', '综合', '—', 76, 70, None, None, '2024', '补长上下文数值推理'],
    ['CNFinBench', '金融', '综合', '—', 71, None, None, None, '2026', '较新，先观察'],
    ['FinanceBench', '金融', '开放问答', '规模小', 63, 55, None, None, '2023', '规模小 → 可直接扩题量'],
    ['HumanEval', '代码', '函数补全', '164', 97, 96, None, None, '2021', '完全饱和 → 已被 SWE-bench 取代'],
    ['SWE-bench Verified', '代码', '真实工程任务', '500', 72, 49, None, None, '2026', '增幅极大，需提难度'],
]
for i, r in enumerate(SAT, 4):
    sat.append(r)
    body(sat, i, 10)
    for c in (2, 3, 4, 5, 6, 7, 8, 9):
        sat.cell(row=i, column=c).alignment = CTR
    for c in (1, 2, 3, 4, 5, 6, 9, 10):
        sat.cell(row=i, column=c).font = BLUE_IN
    sat.cell(row=i, column=5).number_format = '0'
    sat.cell(row=i, column=6).number_format = '0'
    sat.cell(row=i, column=7).value = f'=IF(F{i}="","-",E{i}-F{i})'
    sat.cell(row=i, column=7).number_format = '+0;-0;0'
    sat.cell(row=i, column=8).value = (
        f'=IF(E{i}>=90,"⛔已饱和",'
        f'IF(E{i}>=80,"⚠接近饱和",'
        f'IF(AND(F{i}<>"",E{i}-F{i}>=15),"⚠疑似污染","✓仍有空间")))')
    sat.cell(row=i, column=8).font = Font(name=F, bold=True, size=10)
    sat.cell(row=i, column=8).alignment = CTR

lastsat = len(SAT) + 3
sat.conditional_formatting.add(f'H4:H{lastsat}', FormulaRule(
    formula=['ISNUMBER(SEARCH("已饱和",H4))'],
    fill=PatternFill('solid', bgColor='C00000'), font=Font(color='FFFFFF', bold=True)))
sat.conditional_formatting.add(f'H4:H{lastsat}', FormulaRule(
    formula=['ISNUMBER(SEARCH("⚠",H4))'],
    fill=PatternFill('solid', bgColor='FFC000'), font=Font(bold=True)))
sat.conditional_formatting.add(f'H4:H{lastsat}', FormulaRule(
    formula=['ISNUMBER(SEARCH("仍有空间",H4))'],
    fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100', bold=True)))

rr = lastsat + 2
sat.cell(row=rr, column=1, value='判定规则').font = SEC_FONT
sat.cell(row=rr, column=1).fill = SEC_FILL
RULES = [
    ['SOTA >= 90%', '已饱和 —— 不要再投入，转换题型或提高难度'],
    ['SOTA 80-90%', '接近饱和 —— 只做难题增量，或加对抗样本'],
    ['半年增幅 >= 15pt', '疑似数据污染 —— 机会：做防污染动态更新版本'],
    ['SOTA 30-70%', '理想区间 —— 有 Failure 又有区分度，最值得投入'],
    ['SOTA < 20%', '可能是题目或评分标准有问题，先排查再下结论'],
]
for i, (k, v) in enumerate(RULES, rr + 1):
    sat.cell(row=i, column=1, value=k).font = BOLD
    sat.cell(row=i, column=2, value=v).font = B_FONT
    sat.merge_cells(start_row=i, start_column=2, end_row=i, end_column=10)
    for c in range(1, 11):
        sat.cell(row=i, column=c).border = BD
        sat.cell(row=i, column=c).alignment = WRAP
widths(sat, {'A': 24, 'B': 10, 'C': 16, 'D': 12, 'E': 12, 'F': 14, 'G': 9, 'H': 15, 'I': 12, 'J': 38})

# =====================================================================
# Sheet 7  探针验收表
# =====================================================================
pb = wb.create_sheet('06_探针验收表')
pb['A1'] = '20 题探针实验：8 条 Gate 验收表（立项前最后一关）'
pb['A1'].font = TITLE
pb.merge_cells('A1:H1')
pb['A2'] = '投入：20-30 题（易7/中7/难6）× 6 个梯度模型（强2/中2/弱2）× 2 位专家盲评 + LLM Judge 评 3 遍'
pb['A2'].font = Font(name=F, size=10, italic=True, color='555555')
pb.merge_cells('A2:H2')
pb.append([])
hdr7 = ['Gate', '验收维度', '实测指标', '通过门槛', '实测值', '判定', '不达标对策', '责任人']
pb.append(hdr7)
head(pb, 4, 8)

PB = [
    ['G1', '①真实任务价值', '专家认可「值得考」比例 %', 80, 92, None, '换任务，回到真实工作流重新拆解', '专家组'],
    ['G2', '②明显Failure', 'SOTA 模型平均得分 %', 30, 47, None, '过高→加难题；过低→排查评分是否过严', '算法'],
    ['G3', '③区分度潜力', '最强与最弱模型分数极差 pt', 25, 38, None, '重设难度梯度或评分维度', '算法'],
    ['G3b', '③排序一致性', '与公认能力排序 Spearman', 0.8, 0.89, None, '极差大但排序乱 = 噪声，需降低评分随机性', '算法'],
    ['G4', '④可客观评判', '专家间一致性 Kappa/Alpha', 0.7, 0.76, None, 'Rubric 改写成二元可判定条件', 'PM'],
    ['G4b', '④Judge 稳定性', 'LLM Judge 自一致率 %', 95, 97, None, '固定 judge 版本/温度=0/规范 prompt', '算法'],
    ['G4c', '④人机一致性', 'LLM Judge vs 专家 Kappa', 0.7, 0.72, None, '继续人工评，或先做 judge 对齐微调', '算法'],
    ['G5', '⑤指导迭代', '失败可归因率 %（归入<10类）', 80, 85, None, '评分需拆解到子能力维度而非总分', 'PM'],
    ['G6', '⑥抗污染寿命', '题目公网可检索率 %（越低越好）', 10, 6, None, '改原创题 / 建私有 holdout 集', 'PM'],
    ['G7', '⑦成本可行', '单题总成本（元）', 300, 260, None, '降 Rubric 粒度或缩减目标规模', 'PM'],
    ['G8', '⑧权威性可采信', '评分点有权威来源溯源率 %', 100, 100, None, '补齐依据溯源与权威性分级', 'PM + 专家'],
]
for i, r in enumerate(PB, 5):
    pb.append(r)
    body(pb, i, 8)
    pb.row_dimensions[i].height = 30
    pb.cell(row=i, column=1).font = BOLD
    pb.cell(row=i, column=1).fill = SEC_FILL
    pb.cell(row=i, column=1).alignment = CTR
    for c in (4, 5):
        pb.cell(row=i, column=c).alignment = CTR
    pb.cell(row=i, column=5).font = BLUE_IN

# G2 门槛是下限但也有上限；G6/G7 是越低越好 → 单独处理
for i in range(5, 5 + len(PB)):
    gate = pb.cell(row=i, column=1).value
    if gate in ('G6', 'G7'):
        pb.cell(row=i, column=6).value = f'=IF(E{i}="","待测",IF(E{i}<=D{i},"✓ PASS","✗ FAIL"))'
    elif gate == 'G2':
        pb.cell(row=i, column=6).value = f'=IF(E{i}="","待测",IF(AND(E{i}>=D{i},E{i}<=70),"✓ PASS","✗ FAIL"))'
    else:
        pb.cell(row=i, column=6).value = f'=IF(E{i}="","待测",IF(E{i}>=D{i},"✓ PASS","✗ FAIL"))'
    pb.cell(row=i, column=6).font = Font(name=F, bold=True, size=10)
    pb.cell(row=i, column=6).alignment = CTR

lastpb = 4 + len(PB)
pb.conditional_formatting.add(f'F5:F{lastpb}', FormulaRule(
    formula=['ISNUMBER(SEARCH("PASS",F5))'],
    fill=PatternFill('solid', bgColor='00B050'), font=Font(color='FFFFFF', bold=True)))
pb.conditional_formatting.add(f'F5:F{lastpb}', FormulaRule(
    formula=['ISNUMBER(SEARCH("FAIL",F5))'],
    fill=PatternFill('solid', bgColor='C00000'), font=Font(color='FFFFFF', bold=True)))

rr = lastpb + 2
pb.cell(row=rr, column=1, value='总体结论').font = SEC_FONT
pb.cell(row=rr, column=1).fill = SEC_FILL
pb.cell(row=rr, column=2, value='PASS 数').font = BOLD
pb.cell(row=rr, column=3).value = f'=COUNTIF(F5:F{lastpb},"✓ PASS")'
pb.cell(row=rr, column=4, value='FAIL 数').font = BOLD
pb.cell(row=rr, column=5).value = f'=COUNTIF(F5:F{lastpb},"✗ FAIL")'
pb.cell(row=rr, column=6, value='立项建议').font = BOLD
pb.cell(row=rr, column=7).value = (
    f'=IF(COUNTIF(F5:F{lastpb},"待测")>0,"探针未完成",'
    f'IF(COUNTIF(F5:F{lastpb},"✗ FAIL")=0,"✅ 全部通过，建议立项",'
    f'IF(COUNTIF(F5:F{lastpb},"✗ FAIL")<=2,"⚠ 局部不达标，整改后复测","⛔ 设计存在根本问题，回到选题")))')
pb.cell(row=rr, column=7).font = Font(name=F, bold=True, size=11, color='C00000')
for c in range(1, 9):
    pb.cell(row=rr, column=c).border = BD
    pb.cell(row=rr, column=c).alignment = CTR
pb.merge_cells(start_row=rr, start_column=7, end_row=rr, end_column=8)

rr += 2
pb.cell(row=rr, column=1, value='难度分层设计（保证 benchmark 寿命）').font = SEC_FONT
pb.cell(row=rr, column=1).fill = SEC_FILL
pb.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=8)
for c in range(1, 9):
    pb.cell(row=rr, column=c).fill = SEC_FILL
    pb.cell(row=rr, column=c).border = BD
LV = [
    ['Easy  30%', '区分「弱模型 vs 中等模型」，预期 SOTA 85-95%'],
    ['Medium 40%', '区分「中等 vs 强」——主战场，预期 SOTA 50-70%'],
    ['Hard  30%', '区分「强 vs SOTA」并预留未来 2 年空间，预期 SOTA 15-35%'],
]
for i, (k, v) in enumerate(LV, rr + 1):
    pb.cell(row=i, column=1, value=k).font = BOLD
    pb.cell(row=i, column=2, value=v).font = B_FONT
    pb.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
    for c in range(1, 9):
        pb.cell(row=i, column=c).border = BD
        pb.cell(row=i, column=c).alignment = WRAP
widths(pb, {'A': 8, 'B': 20, 'C': 32, 'D': 13, 'E': 12, 'F': 12, 'G': 44, 'H': 12})

# =====================================================================
# Sheet 8  Idea 卡片模板
# =====================================================================
cd = wb.create_sheet('07_Idea卡片')
cd['A1'] = 'Idea 一页纸卡片（判定为 GO 后填写，作为 PRD 前置文档）'
cd['A1'].font = TITLE
cd.merge_cells('A1:D1')
cd.append([])
CARD = [
    ('SECTION', '一、基本信息'),
    ('ROW', 'Idea 编号 / 名称', ''),
    ('ROW', '垂类领域 / 目标能力', ''),
    ('ROW', '评测范式', ''),
    ('ROW', '来源象限 / 发现渠道', ''),
    ('ROW', '提出人 / 日期', ''),
    ('SECTION', '二、价值论证（对应 Gate ①）'),
    ('ROW', '谁在做这个任务（具体角色）', ''),
    ('ROW', '做得多频繁', ''),
    ('ROW', '做错的后果是什么', ''),
    ('ROW', '谁会使用这个 benchmark 的结论', ''),
    ('ROW', '他们用它做什么决策', ''),
    ('SECTION', '三、Failure 证据（对应 Gate ②③）'),
    ('ROW', '探针实测 SOTA 得分', ''),
    ('ROW', '典型失败案例（3 个）', ''),
    ('ROW', '失败模式归类', ''),
    ('ROW', '模型间分数极差 / 排序一致性', ''),
    ('SECTION', '四、评测设计（对应 Gate ④⑤）'),
    ('ROW', '题型与难度分层（易/中/难）', ''),
    ('ROW', 'Rubric 设计（粒度 / 权重规则）', ''),
    ('ROW', '打分方式（专家 / LLM Judge / 混合）', ''),
    ('ROW', '依据溯源要求（权威性分级标准）', ''),
    ('ROW', '结果如何拆解到子能力（可指导迭代）', ''),
    ('SECTION', '五、数据与合规（对应 Gate ⑥）'),
    ('ROW', '数据来源', ''),
    ('ROW', '版权 / 隐私 / 伦理审批', ''),
    ('ROW', '防污染策略', ''),
    ('ROW', '更新机制与节奏', ''),
    ('SECTION', '六、资源与排期（对应 Gate ⑦）'),
    ('ROW', '专家团队（资质 / 人数）', ''),
    ('ROW', '目标题量 / 单题成本 / 总预算', ''),
    ('ROW', '里程碑排期', ''),
    ('SECTION', '七、权威性建设（对应 Gate ⑧）'),
    ('ROW', '专家背书 / 可公示资质', ''),
    ('ROW', '方法论发布计划（论文 / 技术报告 / 开源）', ''),
    ('ROW', '外部合作与榜单落地', ''),
    ('SECTION', '八、风险与退出'),
    ('ROW', 'Top 3 风险及应对', ''),
    ('ROW', '什么情况下终止项目', ''),
]
r = 3
for item in CARD:
    if item[0] == 'SECTION':
        cd.cell(row=r, column=1, value=item[1]).font = Font(name=F, bold=True, size=11, color='FFFFFF')
        for c in range(1, 4):
            cd.cell(row=r, column=c).fill = SUB_FILL
            cd.cell(row=r, column=c).border = BD
        cd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        cd.row_dimensions[r].height = 24
    else:
        cd.cell(row=r, column=1, value=item[1]).font = BOLD
        cd.cell(row=r, column=1).fill = PatternFill('solid', fgColor='F2F2F2')
        cd.cell(row=r, column=2, value=item[2]).font = BLUE_IN
        cd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        for c in range(1, 4):
            cd.cell(row=r, column=c).border = BD
            cd.cell(row=r, column=c).alignment = WRAP
        cd.row_dimensions[r].height = 32
    r += 1
widths(cd, {'A': 36, 'B': 60, 'C': 40})

# =====================================================================
# Sheet 9  Idea 生产流水线
# =====================================================================
pl = wb.create_sheet('08_生产流水线')
pl['A1'] = 'Idea 生产流水线：价值源 × 难度源 × 设计源 交叉法'
pl['A1'].font = TITLE
pl.merge_cells('A1:F1')
pl['A2'] = ('好 Idea = 价值源 × 难度源 × 设计源 三者交叉。只有价值源 → 伪需求（模型其实已做得很好）；'
            '只有难度源 → 无人在意的失败；只有设计源 → 别人 benchmark 的补丁。')
pl['A2'].font = Font(name=F, size=10, italic=True, color='C00000')
pl.merge_cells('A2:F2')
pl.append([])
hdr9 = ['步骤', '动作', '用哪些源', '输出物', '否决条件（早否早省钱）', '示例：PDF 财报异常识别与证据引用']
pl.append(hdr9)
head(pl, 4, 6, h=38)

PIPE = [
    ['Step 1\n价值定位', '① 问业务方「现在最不知道模型哪项能力到底行不行」\n② 真实用户 Query 聚类，找高频 × 高价值 × 当前表现差的簇',
     'S8 业务方不确定性\nS2 真实用户 Query',
     '锁定一个具体的真实任务\n（含角色 / 频率 / 错误后果）',
     '说不出具体角色、频率或错误后果\n→ 直接否决',
     '业务方：「不知道新模型有没有提升 PDF 财务分析」\nQuery 高频簇 = 读财报→找异常→引用证据→给结论\n（不是 CFA 选择题）'],
    ['Step 2\nFailure 验证', '拿当前 SOTA 跑 20 个真实任务，采集失败案例并归类\n→ 把 Failure 抽象成可评测的能力维度',
     'S3 模型 Failure Case\nS6 论文与新能力',
     'Failure 清单 + 能力抽象结果',
     'SOTA 得分 >85% → 否决（已饱和）\nSOTA <20% → 先排查题目/评分',
     '失败集中在：引用错数字、跨页勾稽错、去年数当今年、口径未调整\n→ 抽象为 Long-context retrieval + Numerical grounding\n+ Cross-page reasoning + Citation correctness'],
    ['Step 3\n区分度验证\n⭐最便宜的关卡', '同一批 20 题跑 GPT / Claude / Gemini / DeepSeek / Qwen\n看模型之间差距与排序是否稳定',
     'S7 竞品模型差异',
     '分数极差 + Spearman 排序一致性',
     '极差 <25pt → 否决（无区分度）\n排序混乱 → 是噪声不是区分度',
     '极差 30pt+，排序与公认能力一致\n→ 区分度通过 ✓'],
    ['Step 4\nRubric 设计', '对专家用「三问法」，不要问「给我出 100 道难题」：\n① 最难的 5 件事？② 新人最容易在哪犯错？③ 什么任务一看就知道专不专业？',
     'S4 专家工作流',
     'case-specific Rubric\n（评分点 + 权重 + 权威依据）',
     '专家给不出可判定的标准\n→ 降级为 WATCH，先做方法论',
     '分析师：「新人最容易错在口径调整和附注勾稽」\n「一眼看专不专业：会不会去翻附注」\n→ 直接成为 Rubric 核心项，且答案唯一可自动判分'],
    ['Step 5\n差异化定位', '研究现有 Benchmark 的 taxonomy / task design / 已披露 failure\n重点不是照搬题，是找它没覆盖、覆盖不深、已饱和的地方',
     'S1 现有 Benchmark\nS5 产品与 Agent 场景',
     '差异化声明（一句话说清独特性）',
     '已有 Bench 完全覆盖且未饱和\n→ 否决',
     'FinanceBench 规模小；CNFinBench 偏单轮 QA；\n两者均无 citation 正确性维度\n→ 差异化 = 任务式 + 引用正确性单独计分 ✓'],
    ['Step 6\n立项', '填 07_Idea卡片 + 跑 06_探针验收表 的 8 条 Gate',
     '—',
     '立项材料（PRD 前置文档）',
     '8 条 Gate 任一 FAIL → 整改后复测\n≥3 条 FAIL → 回到选题',
     '8 条 Gate 预估全过 → GO 立项（见 01_Idea池 ID-16）'],
]
for i, r in enumerate(PIPE, 5):
    pl.append(r)
    body(pl, i, 6)
    pl.row_dimensions[i].height = 92
    pl.cell(row=i, column=1).font = Font(name=F, bold=True, size=10, color='FFFFFF')
    pl.cell(row=i, column=1).fill = SUB_FILL
    pl.cell(row=i, column=1).alignment = CTR
    t = r[2]
    if 'S8' in t or 'S2' in t or 'S4' in t:
        pl.cell(row=i, column=3).fill = VAL
    elif 'S3' in t or 'S7' in t or 'S6' in t:
        pl.cell(row=i, column=3).fill = DIF
    elif 'S1' in t or 'S5' in t:
        pl.cell(row=i, column=3).fill = DES
    pl.cell(row=i, column=3).alignment = CTR
    pl.cell(row=i, column=5).font = Font(name=F, size=10, color='C00000')

r = 5 + len(PIPE) + 1
pl.cell(row=r, column=1, value='Failure → Capability 抽象对照表（Step 2 的核心动作）').font = SEC_FONT
for c in range(1, 7):
    pl.cell(row=r, column=c).fill = SEC_FILL
    pl.cell(row=r, column=c).border = BD
pl.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)

ABS = [
    ['观察到的失败现象', '抽象出的能力维度', '可对接的评测设计'],
    ['长 PDF / 长病历中引用错数字', 'Long-context retrieval + Numerical grounding', '定位题 + 数值精确匹配'],
    ['跨页、跨表、跨期信息整合出错', 'Cross-page / Cross-document reasoning', '勾稽计算题，答案唯一'],
    ['引用了不存在的页码、条文、文献', 'Citation correctness / Attribution', '引用位置可自动校验，单独计分'],
    ['明明不确定却给出确定结论', 'Uncertainty calibration', '要求输出置信度，算校准误差'],
    ['引用了过时版本的指南 / 法条 / 准则', 'Knowledge recency + Version awareness', '同一问题的新旧版本对照题'],
    ['多步任务执行到中途跑偏', 'Agent planning + 任务保持', '轨迹评分（trajectory scoring）'],
    ['被指出错误后仍坚持原答案', 'Self-correction', '两轮交互：给出反证后看是否修正'],
    ['工具/函数参数填错、调用顺序错', 'Tool use', '沙箱环境 + 调用序列与参数校验'],
    ['口径、单位、时间基准未做调整', 'Domain convention grounding', '专家 Rubric 核心项（新人高频错点）'],
]
for i, row in enumerate(ABS, r + 1):
    pl.cell(row=i, column=1, value=row[0])
    pl.cell(row=i, column=2, value=row[1])
    pl.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    pl.cell(row=i, column=5, value=row[2])
    pl.merge_cells(start_row=i, start_column=5, end_row=i, end_column=6)
    for c in range(1, 7):
        x = pl.cell(row=i, column=c)
        x.border = BD
        x.alignment = WRAP
        x.font = BOLD if i == r + 1 else B_FONT
        if i == r + 1:
            x.fill = PatternFill('solid', fgColor='D9E1F2')
    pl.row_dimensions[i].height = 26

r = r + len(ABS) + 2
pl.cell(row=r, column=1, value='5 条纪律（写在墙上）').font = SEC_FONT
for c in range(1, 7):
    pl.cell(row=r, column=c).fill = SEC_FILL
    pl.cell(row=r, column=c).border = BD
pl.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)

RULE5 = [
    ['1. Step 3 尽早做', '20 题 × 4 个模型成本极低，却能直接否掉「无区分度」的 Idea。区分度验证放最前面，能省掉后面 90% 的白干。'],
    ['2. 别问专家出难题', '「给我出 100 道难题」得到的是偏题怪题。要问三问：最难的 5 件事 / 新人易错点 / 一眼看出专不专业的任务。'],
    ['3. 研究 Bench 不抄题', '看 GAIA、HLE、BrowseComp、SWE-bench 是为了学 taxonomy、task design 和它披露的 failure case，不是搬题目。'],
    ['4. Failure 必须抽象成能力', '「模型引用错数字」不是能力，「Numerical grounding」才是。不抽象就无法成体系、无法指导迭代。'],
    ['5. 没区分度的题价值有限', '真实价值再高，如果所有模型都做得一样好或一样差，这个 benchmark 不产生信息量。'],
]
for i, (k, v) in enumerate(RULE5, r + 1):
    pl.cell(row=i, column=1, value=k).font = BOLD
    pl.cell(row=i, column=2, value=v).font = B_FONT
    pl.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    for c in range(1, 7):
        pl.cell(row=i, column=c).border = BD
        pl.cell(row=i, column=c).alignment = WRAP
    pl.row_dimensions[i].height = 34

widths(pl, {'A': 15, 'B': 46, 'C': 22, 'D': 30, 'E': 32, 'F': 58})
pl.freeze_panes = 'A5'

out = r'C:\Users\vegavjzhang\Desktop\Benchmark Idea Radar\Benchmark_Idea_Radar.xlsx'
wb.save(out)
print('SAVED:', out)
print('sheets:', wb.sheetnames)
